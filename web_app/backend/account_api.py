"""Authentication, engineer settings, artifact, and administrator APIs."""

from __future__ import annotations

import os
import secrets
import csv
from io import BytesIO, StringIO
from typing import Any, Dict, Optional

from fastapi import APIRouter, Body, Depends, File, Header, HTTPException, Request, Response, UploadFile
from pydantic import BaseModel, Field
from starlette.middleware.base import BaseHTTPMiddleware
from starlette.responses import JSONResponse

from web_app.backend.identity_store import (
    CurrentUser,
    IdentityStore,
    canonical_company_email,
)
from web_app.backend.admin_log_service import get_operation_log, list_operation_logs


SESSION_COOKIE = "cad_session"
DOMAIN_COOKIE = "cad_email_domain"
AUTH_EXEMPT_PREFIXES = (
    "/api/health",
    "/api/auth/login",
    "/api/auth/config",
)
ADMIN_ONLY_EXACT_PATHS = {
    "/api/examples",
    "/api/processed/fan-20260625",
    "/api/tolerance/cases",
    "/api/tolerance/open-local",
    "/tolerance-inspector.html",
    "/tolerance-audit-v2.html",
}
ADMIN_ONLY_PREFIXES = (
    "/api/examples/",
    "/api/tolerance/debug/",
    "/api/tolerance/audit/",
    "/api/tolerance/download/",
)


store = IdentityStore()
router = APIRouter(prefix="/api", tags=["accounts"])


def _request_token(request: Request) -> Optional[str]:
    authorization = request.headers.get("authorization", "")
    if authorization.lower().startswith("bearer "):
        return authorization[7:].strip()
    return request.cookies.get(SESSION_COOKIE)


def require_user(request: Request) -> CurrentUser:
    user = getattr(request.state, "current_user", None)
    if not user:
        raise HTTPException(status_code=401, detail="Authentication required")
    return user


def require_admin(user: CurrentUser = Depends(require_user)) -> CurrentUser:
    if not user.is_admin:
        raise HTTPException(status_code=403, detail="Administrator permission required")
    return user


def is_admin_only_path(path: str) -> bool:
    return path in ADMIN_ONLY_EXACT_PATHS or any(
        path.startswith(prefix) for prefix in ADMIN_ONLY_PREFIXES
    )


def assert_model_access(user: CurrentUser, model_id: str) -> None:
    if not store.can_access_model(user, model_id):
        # Return 404 to avoid leaking another engineer's model identifiers.
        raise HTTPException(status_code=404, detail="Model not found")


def require_email_delivery_key(
    x_api_key: Optional[str] = Header(default=None, alias="X-API-Key"),
) -> None:
    configured = os.environ.get("CAD_EMAIL_DELIVERY_API_KEY", "")
    if not configured:
        raise HTTPException(status_code=503, detail="Email delivery integration is not configured")
    if not x_api_key or not secrets.compare_digest(x_api_key, configured):
        raise HTTPException(status_code=401, detail="Invalid email delivery API key")


class AuthenticationMiddleware(BaseHTTPMiddleware):
    """Require a revocable session for API and protected generated files."""

    async def dispatch(self, request: Request, call_next):
        path = request.url.path
        token = _request_token(request)
        user = store.user_for_token(token)
        request.state.current_user = user

        is_api = path == "/api" or path.startswith("/api/")
        exempt = any(path.startswith(prefix) for prefix in AUTH_EXEMPT_PREFIXES)
        configured_api_key = os.environ.get("CAD_EXTERNAL_PREDICTION_API_KEY", "")
        external_api_authorized = (
            path.startswith("/api/tolerance/external-predictions")
            and bool(configured_api_key)
            and secrets.compare_digest(
                request.headers.get("X-API-Key", ""), configured_api_key
            )
        )
        email_api_key = os.environ.get("CAD_EMAIL_DELIVERY_API_KEY", "")
        email_integration_authorized = (
            path.startswith("/api/integrations/email/credential-notifications")
            and bool(email_api_key)
            and secrets.compare_digest(request.headers.get("X-API-Key", ""), email_api_key)
        )
        if is_api and not exempt and not user and not external_api_authorized and not email_integration_authorized:
            return JSONResponse({"detail": "Authentication required"}, status_code=401)

        password_change_paths = {
            "/api/auth/login", "/api/auth/me", "/api/auth/password", "/api/auth/logout", "/api/auth/config",
        }
        if user and not user.is_admin and user.must_change_password and is_api and path not in password_change_paths:
            return JSONResponse({"detail": "PASSWORD_CHANGE_REQUIRED"}, status_code=403)

        admin_only_reference = is_admin_only_path(path)
        if admin_only_reference and user and not user.is_admin:
            return JSONResponse(
                {"detail": "Administrator permission required"},
                status_code=403,
            )

        if user and path.startswith("/api/files/"):
            remainder = path.removeprefix("/api/files/")
            model_id = remainder.split("/", 1)[0]
            if model_id and not store.can_access_model(user, model_id):
                return JSONResponse({"detail": "File not found"}, status_code=404)

        return await call_next(request)


class LoginRequest(BaseModel):
    email_local: Optional[str] = Field(default=None, min_length=1, max_length=64)
    email_domain: Optional[str] = Field(default=None, max_length=120)
    username: Optional[str] = Field(default=None, min_length=1, max_length=80)
    password: str = Field(min_length=1, max_length=256)


class PasswordChangeRequest(BaseModel):
    current_password: str
    new_password: str = Field(min_length=10, max_length=256)
    confirm_password: str = Field(min_length=10, max_length=256)


class UserCreateRequest(BaseModel):
    display_name: str = Field(min_length=1, max_length=120)
    email: str = Field(min_length=3, max_length=80)
    role: str = "ENGINEER"


class UserBulkDeleteRequest(BaseModel):
    user_ids: list[str] = Field(min_length=1, max_length=200)


class PasswordResetRequest(BaseModel):
    new_password: Optional[str] = Field(default=None, min_length=10, max_length=256)


class NotificationAckRequest(BaseModel):
    status: str
    provider_message_id: Optional[str] = None
    error: Optional[str] = None


class RetentionRequest(BaseModel):
    days: int = Field(ge=1, le=3650)


class CadConcurrencyRequest(BaseModel):
    analysis_max_concurrency: int = Field(ge=1, le=8)
    annotation_max_concurrency: int = Field(ge=1, le=8)


class CompanyEmailDomainRequest(BaseModel):
    domain: str = Field(min_length=3, max_length=120)


class ImportDomainResolutionRequest(BaseModel):
    rows: list[Dict[str, Any]] = Field(default_factory=list)
    resolutions: Dict[str, str] = Field(default_factory=dict)


class RecommendationTagRequest(BaseModel):
    key: str = Field(min_length=1, max_length=160)
    name: str = Field(min_length=1, max_length=160)
    dimension: str = Field(default="CUSTOM", max_length=40)
    description: Optional[str] = None


class TagAssignmentRequest(BaseModel):
    tag_ids: list[str] = Field(default_factory=list)


@router.post("/auth/login")
def login(payload: LoginRequest, request: Request, response: Response):
    identifier = payload.email_local or payload.username or ""
    domains = store.list_company_email_domains()
    selected_domain = payload.email_domain or domains[0]
    try:
        store.canonical_company_email(identifier, selected_domain)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    user = store.authenticate(identifier, payload.password, selected_domain)
    if not user:
        raise HTTPException(status_code=401, detail="Invalid company email or password")
    client_ip = request.client.host if request.client else None
    token, expires_at = store.create_session(
        user.id,
        client_ip=client_ip,
        user_agent=request.headers.get("user-agent"),
    )
    response.set_cookie(
        SESSION_COOKIE,
        token,
        httponly=True,
        secure=os.environ.get("CAD_COOKIE_SECURE", "0") == "1",
        samesite="strict",
        max_age=12 * 60 * 60,
        path="/",
    )
    response.set_cookie(
        DOMAIN_COOKIE,
        user.email.rsplit("@", 1)[1],
        httponly=False,
        secure=os.environ.get("CAD_COOKIE_SECURE", "0") == "1",
        samesite="strict",
        max_age=365 * 24 * 60 * 60,
        path="/",
    )
    return {"status": "ok", "user": user.as_dict(), "expires_at": expires_at.isoformat()}


@router.get("/auth/config")
def auth_config():
    return {
        "status": "ok",
        "company_email_domains": store.list_company_email_domains(),
        "domain_cookie": DOMAIN_COOKIE,
    }


@router.post("/auth/logout")
def logout(request: Request, response: Response, user: CurrentUser = Depends(require_user)):
    store.revoke_session(_request_token(request))
    response.delete_cookie(SESSION_COOKIE, path="/")
    return {"status": "ok"}


@router.get("/auth/me")
def me(user: CurrentUser = Depends(require_user)):
    return {"status": "ok", "user": user.as_dict()}


@router.put("/auth/password")
def change_password(
    payload: PasswordChangeRequest,
    request: Request,
    response: Response,
    user: CurrentUser = Depends(require_user),
):
    if payload.new_password != payload.confirm_password:
        raise HTTPException(status_code=400, detail="New password confirmation does not match")
    if not user.is_admin and not user.must_change_password:
        raise HTTPException(
            status_code=403,
            detail="Password changes for engineers must be performed by an administrator",
        )
    try:
        store.change_password(user.id, payload.current_password, payload.new_password)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    store.revoke_session(_request_token(request))
    response.delete_cookie(SESSION_COOKIE, path="/")
    return {"status": "ok", "reauthentication_required": True}


@router.get("/engineer/preferences")
def get_preferences(user: CurrentUser = Depends(require_user)):
    return {"status": "ok", "preferences": store.get_preferences(user.id)}


@router.put("/engineer/preferences")
def put_preferences(
    payload: Dict[str, Any] = Body(...),
    user: CurrentUser = Depends(require_user),
):
    try:
        preferences = store.update_preferences(user.id, payload)
    except (TypeError, ValueError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return {"status": "ok", "preferences": preferences}


@router.delete("/engineer/preferences")
def reset_preferences(user: CurrentUser = Depends(require_user)):
    return {
        "status": "ok",
        "preferences": store.reset_preferences(user.id),
    }


@router.get("/recommendation-tags")
def list_recommendation_tags(user: CurrentUser = Depends(require_user)):
    return {"status": "ok", "tags": store.list_recommendation_tags()}


@router.get("/engineer/artifacts")
def list_artifacts(
    owner_user_id: Optional[str] = None,
    user: CurrentUser = Depends(require_user),
):
    return {"status": "ok", "artifacts": store.list_artifacts(user, owner_user_id)}


@router.delete("/engineer/artifacts/{artifact_id}")
def delete_artifact(artifact_id: str, user: CurrentUser = Depends(require_user)):
    try:
        found = store.delete_artifact(user, artifact_id)
    except PermissionError as exc:
        raise HTTPException(status_code=403, detail=str(exc)) from exc
    if not found:
        raise HTTPException(status_code=404, detail="Artifact not found")
    return {"status": "ok"}


@router.get("/engineer/tolerance-cases")
def list_personal_tolerance_cases(user: CurrentUser = Depends(require_user)):
    return {"status": "ok", "cases": store.personal_cases(user.id)}


@router.delete("/engineer/tolerance-cases")
def clear_personal_tolerance_cases(user: CurrentUser = Depends(require_user)):
    deleted_count = store.clear_personal_cases(user.id)
    return {"status": "ok", "deleted_count": deleted_count}


@router.delete("/engineer/tolerance-cases/{case_record_id}")
def delete_personal_tolerance_case(
    case_record_id: str,
    user: CurrentUser = Depends(require_user),
):
    if not store.delete_personal_case(user.id, case_record_id):
        raise HTTPException(status_code=404, detail="Personal tolerance case not found")
    return {"status": "ok"}


@router.get("/engineer/trash")
def list_engineer_trash(user: CurrentUser = Depends(require_user)):
    return {
        "status": "ok",
        "retention_days": store.trash_retention_days(),
        "records": store.list_trash(user),
    }


@router.post("/engineer/trash/{trash_id}/restore")
def restore_engineer_trash(trash_id: str, user: CurrentUser = Depends(require_user)):
    try:
        result = store.restore_trash(user, trash_id)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except PermissionError as exc:
        raise HTTPException(status_code=403, detail=str(exc)) from exc
    except (ValueError, TypeError) as exc:
        raise HTTPException(status_code=409, detail=f"Cannot restore record: {exc}") from exc
    return {"status": "ok", **result}


@router.get("/admin/users")
def admin_list_users(admin: CurrentUser = Depends(require_admin)):
    return {"status": "ok", "users": store.list_users(include_initial_passwords=True)}


@router.post("/admin/users", status_code=201)
def admin_create_user(
    payload: UserCreateRequest,
    admin: CurrentUser = Depends(require_admin),
):
    try:
        user = store.create_user(payload.model_dump(), admin.id)
    except ValueError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc
    return {"status": "ok", "user": user}


@router.delete("/admin/users/{user_id}")
def admin_delete_user(user_id: str, admin: CurrentUser = Depends(require_admin)):
    try:
        deleted = store.delete_users([user_id], admin.id)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return {"status": "ok", "deleted_count": 1, "users": deleted}


@router.post("/admin/users/bulk-delete")
def admin_bulk_delete_users(
    payload: UserBulkDeleteRequest,
    admin: CurrentUser = Depends(require_admin),
):
    try:
        deleted = store.delete_users(payload.user_ids, admin.id)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return {"status": "ok", "deleted_count": len(deleted), "users": deleted}


@router.patch("/admin/users/{user_id}")
def admin_update_user(
    user_id: str,
    payload: Dict[str, Any] = Body(...),
    admin: CurrentUser = Depends(require_admin),
):
    if user_id == admin.id and payload.get("is_active") is False:
        raise HTTPException(status_code=400, detail="Administrator cannot disable the active account")
    try:
        user = store.update_user(user_id, payload, admin.id)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return {"status": "ok", "user": user}


@router.post("/admin/users/{user_id}/reset-password")
def admin_reset_password(
    user_id: str,
    payload: PasswordResetRequest,
    admin: CurrentUser = Depends(require_admin),
):
    try:
        initial_password = store.reset_password(user_id, payload.new_password, admin.id)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    target = next((item for item in store.list_users() if item["id"] == user_id), None)
    return {
        "status": "ok",
        "must_change_password": bool(target and target["must_change_password"]),
        "initial_password": initial_password,
    }


def _import_rows_with_domain_review(
    rows: list[Dict[str, Any]], actor_user_id: str
) -> tuple[list[Dict[str, Any]], list[Dict[str, Any]], list[Dict[str, Any]]]:
    """Create valid rows while retaining unknown-domain rows for an admin decision."""
    allowed_domains = store.list_company_email_domains()
    created: list[Dict[str, Any]] = []
    unresolved: list[Dict[str, Any]] = []
    rejected: list[Dict[str, Any]] = []

    for offset, row in enumerate(rows, start=2):
        row_number = row.get("row_number") or offset
        raw_email = str(row.get("email") or row.get("account") or "").strip().lower()
        display_name = str(row.get("display_name") or row.get("name") or "").strip()
        role = str(row.get("role") or "ENGINEER").strip().upper()
        try:
            if not display_name:
                raise ValueError("Display name is required")
            if role not in {"ADMIN", "ENGINEER"}:
                raise ValueError(f"Invalid role: {role}")

            if "@" in raw_email:
                local_part, source_domain = raw_email.rsplit("@", 1)
                source_domain = store._normalize_domain(source_domain)
            else:
                local_part = raw_email
                source_domain = allowed_domains[0]

            if source_domain not in allowed_domains:
                # Validate the local part and total identifier length without accepting
                # the domain into company policy yet.
                canonical_company_email(
                    f"{local_part}@{source_domain}",
                    allowed_domains=[*allowed_domains, source_domain],
                )
                unresolved.append({
                    "row_number": row_number,
                    "email": f"{local_part}@{source_domain}",
                    "local_part": local_part,
                    "source_domain": source_domain,
                    "display_name": display_name,
                    "role": role,
                })
                continue

            normalized_email = canonical_company_email(
                f"{local_part}@{source_domain}", allowed_domains=allowed_domains
            )
            created.append(store.create_user({
                "email": normalized_email,
                "display_name": display_name,
                "role": role,
            }, actor_user_id))
        except (ValueError, TypeError) as exc:
            rejected.append({
                "row_number": row_number,
                "email": raw_email,
                "display_name": display_name,
                "reason": str(exc),
            })

    if not created and not unresolved and not rejected:
        raise ValueError("The import file contains no account rows")
    return created, unresolved, rejected


@router.post("/admin/users/import", status_code=201)
async def admin_import_users(
    file: UploadFile = File(...),
    admin: CurrentUser = Depends(require_admin),
):
    filename = (file.filename or "").lower()
    content = await file.read()
    try:
        if filename.endswith(".xlsx"):
            from openpyxl import load_workbook
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
            sheet = workbook.active
            values = list(sheet.iter_rows(values_only=True))
            if not values:
                raise ValueError("Excel file is empty")
            headers = [str(value or "").strip().lower() for value in values[0]]
            rows = [dict(zip(headers, row)) for row in values[1:] if any(value is not None for value in row)]
        elif filename.endswith(".csv"):
            decoded = content.decode("utf-8-sig")
            rows = list(csv.DictReader(StringIO(decoded)))
        else:
            raise ValueError("Only .xlsx and .csv files are supported")
        aliases = {
            "姓名": "display_name", "名稱": "display_name", "name": "display_name",
            "電子郵件": "email", "信箱": "email", "mail": "email",
            "角色": "role",
        }
        normalized_rows = []
        for row in rows:
            normalized = {}
            for key, value in row.items():
                canonical_key = aliases.get(str(key).strip().lower(), str(key).strip().lower())
                normalized[canonical_key] = value
            normalized_rows.append(normalized)
        created, unresolved, rejected = _import_rows_with_domain_review(
            normalized_rows, admin.id
        )
    except (ValueError, TypeError, UnicodeDecodeError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return {
        "status": "review_required" if unresolved else ("partial" if rejected else "ok"),
        "created_count": len(created),
        "users": created,
        "unresolved_domain_rows": unresolved,
        "rejected_rows": rejected,
    }


@router.post("/admin/users/import/resolve-domains", status_code=201)
def admin_resolve_import_domains(
    payload: ImportDomainResolutionRequest,
    admin: CurrentUser = Depends(require_admin),
):
    current_domains = store.list_company_email_domains()
    normalized_resolutions: Dict[str, str] = {}
    domains_to_add: list[str] = []
    try:
        if not payload.rows:
            raise ValueError("No pending import rows were provided")
        for source, selection in payload.resolutions.items():
            source_domain = store._normalize_domain(source)
            selected = str(selection or "").strip().lower().lstrip("@")
            if selected == "__add__":
                normalized_resolutions[source_domain] = source_domain
                if source_domain not in current_domains and source_domain not in domains_to_add:
                    domains_to_add.append(source_domain)
            elif selected in current_domains:
                normalized_resolutions[source_domain] = selected
            else:
                raise ValueError(f"A resolution is required for @{source_domain}")

        pending_sources = {
            str(row.get("source_domain") or "").strip().lower().lstrip("@")
            for row in payload.rows
        }
        missing = sorted(source for source in pending_sources if source not in normalized_resolutions)
        if missing:
            raise ValueError(f"Missing domain resolutions: {', '.join('@' + item for item in missing)}")

        for domain in domains_to_add:
            store.create_company_email_domain(domain, admin.id)

        resolved_rows = []
        for row in payload.rows:
            source_domain = str(row.get("source_domain") or "").strip().lower().lstrip("@")
            target_domain = normalized_resolutions[source_domain]
            resolved_rows.append({
                "row_number": row.get("row_number"),
                "email": f"{str(row.get('local_part') or '').strip().lower()}@{target_domain}",
                "display_name": row.get("display_name"),
                "role": row.get("role"),
            })
        created, unresolved, rejected = _import_rows_with_domain_review(
            resolved_rows, admin.id
        )
    except (ValueError, TypeError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    if unresolved:
        raise HTTPException(status_code=409, detail="Domain resolution did not resolve every row")
    return {
        "status": "ok",
        "created_count": len(created),
        "users": created,
        "rejected_rows": rejected,
        "domains_added": domains_to_add,
        "company_email_domains": store.list_company_email_domains(),
    }


@router.get("/admin/users/import-template")
def admin_download_user_import_template(
    format: str = "xlsx",
    admin: CurrentUser = Depends(require_admin),
):
    selected_format = str(format).strip().lower()
    rows = [
        ["email", "display_name", "role"],
        [f"engineer@{store.list_company_email_domains()[0]}", "Example Engineer", "ENGINEER"],
    ]
    if selected_format == "csv":
        output = StringIO()
        writer = csv.writer(output, lineterminator="\n")
        writer.writerows(rows)
        content = "\ufeff" + output.getvalue()
        return Response(
            content=content.encode("utf-8"),
            media_type="text/csv; charset=utf-8",
            headers={"Content-Disposition": 'attachment; filename="account-import-template.csv"'},
        )
    if selected_format == "xlsx":
        from openpyxl import Workbook
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "accounts"
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        sheet.column_dimensions["A"].width = 34
        sheet.column_dimensions["B"].width = 24
        sheet.column_dimensions["C"].width = 14
        output = BytesIO()
        workbook.save(output)
        return Response(
            content=output.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": 'attachment; filename="account-import-template.xlsx"'},
        )
    raise HTTPException(status_code=400, detail="Template format must be xlsx or csv")


@router.get("/admin/company-email-domains")
def admin_list_company_email_domains(admin: CurrentUser = Depends(require_admin)):
    return {"status": "ok", "domains": store.list_company_email_domains()}


@router.post("/admin/company-email-domains", status_code=201)
def admin_create_company_email_domain(
    payload: CompanyEmailDomainRequest,
    admin: CurrentUser = Depends(require_admin),
):
    try:
        domain = store.create_company_email_domain(payload.domain, admin.id)
    except ValueError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc
    return {"status": "ok", "domain": domain}


@router.put("/admin/company-email-domains/{domain}")
def admin_update_company_email_domain(
    domain: str,
    payload: CompanyEmailDomainRequest,
    admin: CurrentUser = Depends(require_admin),
):
    try:
        updated = store.update_company_email_domain(domain, payload.domain, admin.id)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc
    return {"status": "ok", "domain": updated}


@router.delete("/admin/company-email-domains/{domain}")
def admin_delete_company_email_domain(
    domain: str,
    admin: CurrentUser = Depends(require_admin),
):
    try:
        store.delete_company_email_domain(domain, admin.id)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc
    return {"status": "ok"}


@router.get("/admin/audit-events")
def admin_audit_events(limit: int = 200, admin: CurrentUser = Depends(require_admin)):
    return {"status": "ok", "events": store.list_audit_events(limit)}


@router.get("/admin/operation-logs")
def admin_operation_logs(
    limit: int = 200,
    status: Optional[str] = None,
    q: Optional[str] = None,
    admin: CurrentUser = Depends(require_admin),
):
    return {
        "status": "ok",
        "logs": list_operation_logs(limit=limit, status=status, query=q),
    }


@router.get("/admin/operation-logs/{operation_id}")
def admin_operation_log_detail(
    operation_id: str,
    admin: CurrentUser = Depends(require_admin),
):
    log = get_operation_log(operation_id)
    if not log:
        raise HTTPException(status_code=404, detail="Operation log not found")
    return {"status": "ok", "log": log}


@router.get("/admin/trash")
def admin_list_trash(
    owner_user_id: Optional[str] = None,
    admin: CurrentUser = Depends(require_admin),
):
    return {
        "status": "ok",
        "retention_days": store.trash_retention_days(),
        "records": store.list_trash(admin, owner_user_id),
    }


@router.post("/admin/trash/{trash_id}/restore")
def admin_restore_trash(trash_id: str, admin: CurrentUser = Depends(require_admin)):
    try:
        result = store.restore_trash(admin, trash_id)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except (ValueError, TypeError) as exc:
        raise HTTPException(status_code=409, detail=f"Cannot restore record: {exc}") from exc
    return {"status": "ok", **result}


@router.delete("/admin/trash/{trash_id}")
def admin_purge_trash(trash_id: str, admin: CurrentUser = Depends(require_admin)):
    if not store.purge_trash(admin, trash_id):
        raise HTTPException(status_code=404, detail="Trash record not found")
    return {"status": "ok", "permanently_deleted": True}


@router.get("/admin/settings/trash-retention")
def admin_get_trash_retention(admin: CurrentUser = Depends(require_admin)):
    return {"status": "ok", "days": store.trash_retention_days()}


@router.put("/admin/settings/trash-retention")
def admin_set_trash_retention(
    payload: RetentionRequest,
    admin: CurrentUser = Depends(require_admin),
):
    return {"status": "ok", "days": store.set_trash_retention_days(payload.days, admin.id)}


@router.get("/admin/settings/cad-concurrency")
def admin_get_cad_concurrency(admin: CurrentUser = Depends(require_admin)):
    return {"status": "ok", **store.cad_concurrency_limits()}


@router.put("/admin/settings/cad-concurrency")
def admin_set_cad_concurrency(
    payload: CadConcurrencyRequest,
    admin: CurrentUser = Depends(require_admin),
):
    try:
        settings = store.set_cad_concurrency_limits(
            payload.analysis_max_concurrency,
            payload.annotation_max_concurrency,
            admin.id,
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return {"status": "ok", **settings}


@router.post("/admin/trash/purge-expired")
def admin_purge_expired_trash(admin: CurrentUser = Depends(require_admin)):
    return {"status": "ok", "purged_count": store.purge_expired_trash(admin.id)}


@router.get("/admin/recommendation-tags")
def admin_list_recommendation_tags(admin: CurrentUser = Depends(require_admin)):
    return {"status": "ok", "tags": store.list_recommendation_tags(include_inactive=True)}


@router.post("/admin/recommendation-tags", status_code=201)
def admin_create_recommendation_tag(
    payload: RecommendationTagRequest,
    admin: CurrentUser = Depends(require_admin),
):
    try:
        tag = store.create_recommendation_tag(payload.model_dump(), admin.id)
    except ValueError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc
    return {"status": "ok", "tag": tag}


@router.patch("/admin/recommendation-tags/{tag_id}")
def admin_update_recommendation_tag(
    tag_id: str,
    payload: Dict[str, Any] = Body(...),
    admin: CurrentUser = Depends(require_admin),
):
    try:
        tag = store.update_recommendation_tag(tag_id, payload, admin.id)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    return {"status": "ok", "tag": tag}


@router.get("/admin/company-cases/tags")
def admin_list_company_case_tags(admin: CurrentUser = Depends(require_admin)):
    return {"status": "ok", "assignments": store.company_case_tags()}


@router.put("/admin/company-cases/{case_id}/tags")
def admin_assign_company_case_tags(
    case_id: str,
    payload: TagAssignmentRequest,
    admin: CurrentUser = Depends(require_admin),
):
    try:
        tag_ids = store.assign_company_case_tags(case_id, payload.tag_ids, admin.id)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return {"status": "ok", "case_id": case_id, "tag_ids": tag_ids}


@router.get("/integrations/email/credential-notifications")
def integration_list_credential_notifications(
    limit: int = 100,
    _: None = Depends(require_email_delivery_key),
):
    return {
        "status": "ok",
        "notifications": store.pending_credential_notifications(limit),
    }


@router.post("/integrations/email/credential-notifications/{notification_id}/ack")
def integration_ack_credential_notification(
    notification_id: str,
    payload: NotificationAckRequest,
    _: None = Depends(require_email_delivery_key),
):
    try:
        found = store.acknowledge_credential_notification(
            notification_id,
            payload.status,
            payload.provider_message_id,
            payload.error,
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    if not found:
        raise HTTPException(status_code=404, detail="Notification not found")
    return {"status": "ok"}
